import asyncio
import io
import json
import os
import stat
from pathlib import Path

import httpx
import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from ccgp_crawler.ai_client import AIModelClient, AIModelError
from ccgp_crawler.ai_config import AIConfigUpdate
from ccgp_crawler.api import create_app
from ccgp_crawler.bid_documents import extract_documents
from ccgp_crawler.bid_models import BidCheckError, BidDecision, CompanyProfile, parse_json_output, validate_decision
from ccgp_crawler.bid_service import NoticeReader
from ccgp_crawler.bid_store import BidTaskStore
from ccgp_crawler.local_store import save_notices
from ccgp_crawler.models import Notice

ANNOUNCEMENT = "二、申请人的资格要求：供应商必须具备有效的建筑施工许可证。满足政府采购法第二十二条。评标时具有本地业绩加分。"
TEXT = "测试公司从事设备销售。公司目前没有建筑施工许可证。"
PROFILE = {"company_name": "测试公司", "facts": [{"id": "f1", "statement": "公司没有建筑施工许可证", "source_id": "F1:company.txt", "quote": "公司目前没有建筑施工许可证。"}], "limitations": ["税收记录资料未提及"]}


def decision(verdict="exclude", assessment="conflict", *, quote="公司目前没有建筑施工许可证。", mandatory=True, reference="供应商必须具备有效的建筑施工许可证。"):
    return {"verdict": verdict, "summary": "资格初筛说明", "pending_checks": [], "requirements": [{
        "requirement": "有效建筑施工许可证", "mandatory": mandatory, "announcement_quote": reference,
        "assessment": assessment, "company_fact_ids": ["f1"] if quote else [], "company_quote": quote, "reason": "结合证据判断",
    }]}


@pytest.fixture
def app(tmp_path):
    app = create_app(database_path=tmp_path / "notices.sqlite3", runtime_dir=tmp_path / "runtime", start_background_services=False)
    app.state.ai_config_store.save(AIConfigUpdate(base_url="https://model.example/v1", model="gpt-test", api_key="test-secret"))
    app.state.bid_manager.interval = 0
    app.state.bid_manager.retry_delays = (0, 0)
    return app


def seed(app, count=8, region="370000", title="设备采购"):
    with app.state.store.connect() as db:
        save_notices(db, [Notice(
            canonical_url=f"https://www.ccgp.gov.cn/cggg/dfgg/gkzb/202610/t{region}_{i}.htm",
            title=f"{title} {i}", published_at="2026-10-07T10:00:00+08:00", region_code=region,
            region_raw="山东", category="dfgg", source_list_url="https://search.ccgp.gov.cn/bxsearch",
        ) for i in range(count)])


def create(client, **fields):
    return client.post("/api/bid-check/tasks", data={"region_code": "370000", "start_date": "2026-10-07", "end_date": "2026-10-07", "keyword": "设备", **fields}, files=[("files", ("company.txt", TEXT.encode(), "text/plain"))])


class FakeAI:
    def __init__(self, outputs=None):
        self.calls = []
        self.outputs = outputs or []

    async def generate(self, config, input_text, **kwargs):
        self.calls.append((json.loads(input_text), kwargs))
        if "sources" in self.calls[-1][0]:
            return json.dumps(PROFILE, ensure_ascii=False)
        output = self.outputs.pop(0) if self.outputs else decision("needs_review", "unknown", quote="")
        if isinstance(output, Exception):
            raise output
        return json.dumps(output, ensure_ascii=False)


class FakeReader:
    async def read(self, url):
        return ANNOUNCEMENT


def process(app, task_id, ai=None):
    # These pre-protocol-2 regression fixtures exercise an existing historical task.
    # New-task snapshots/outputs are covered separately in test_bid_screening.py.
    with app.state.bid_store.connect() as db:
        db.execute('UPDATE tasks SET rules_snapshot=NULL WHERE id=?', (task_id,))
    app.state.ai_client = ai or FakeAI()
    app.state.bid_manager.reader = FakeReader()
    asyncio.run(app.state.bid_manager.process_task(task_id))


def test_all_results_snapshot_not_current_page_and_file_permissions(app):
    seed(app)
    seed(app, count=2, region="440000")
    with TestClient(app) as client:
        assert len(client.get("/api/v1/notices", params={"region_code": "370000"}).json()["items"]) == 6
        draft = create(client).json()
        task_id = draft["id"]
        assert draft["total"] == 8 and draft["estimated_calls"] == 9
        stored = app.state.bid_store.raw(task_id)
        assert "api_key" not in stored["model_config"]
        file = app.state.bid_store.root / task_id / "001.txt"
        assert stat.S_IMODE(file.stat().st_mode) == 0o600
        assert client.get(f"/api/bid-check/tasks/{task_id}").json()["status"] == "draft"
        seed(app, count=1, title="新设备")
        assert client.get(f"/api/bid-check/tasks/{task_id}").json()["total"] == 8
        assert client.post(f"/api/bid-check/tasks/{task_id}/stop").json()["status"] == "stopped"
        assert not file.exists()


def test_independent_judgment_retains_unknown_hides_conflict_and_survives_reload(app):
    seed(app, count=3)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        ai = FakeAI([decision(), decision("needs_review", "unknown", quote=""), decision("exclude", "conflict", quote="invented")])
        process(app, task_id, ai)
        result = client.get(f"/api/bid-check/tasks/{task_id}").json()
        assert result["status"] == "completed" and result["completed"] == 3
        assert result["excluded"] == 1 and result["retained"] == result["needs_review"] == 2
        assert len(result["items"]) == 2
        assert all(item["decision"]["verdict"] == "needs_review" for item in result["items"])
        assert len(ai.calls) == 4 and all(call[1]["timeout"] == 120 for call in ai.calls)
        assert "notice" in ai.calls[1][0] and "canonical_url" in ai.calls[1][0]["notice"]
        assert not (app.state.bid_store.root / task_id).exists()
        assert client.get("/api/bid-check/tasks").json()["items"][0]["id"] == task_id
        store = BidTaskStore(app.state.bid_store.root.parent)
        store.initialize()
        assert store.detail(task_id)["retained"] == 2
        assert client.get(f"/api/bid-check/tasks/{task_id}?page_size=1").json()["has_next"]


@pytest.mark.parametrize("output,expected", [
    (decision("exclude", "unknown", quote=""), "needs_review"),
    (decision("exclude", "conflict", quote=""), "needs_review"),
    (decision("exclude", "conflict", mandatory=False), "retain"),
    (decision("exclude", "conflict", quote="imaginary"), "needs_review"),
    (decision("retain", "unknown", quote=""), "needs_review"),
    (decision(), "exclude"),
])
def test_only_evidenced_mandatory_conflict_excludes(output, expected):
    result = validate_decision(BidDecision.model_validate(output), CompanyProfile.model_validate(PROFILE), ANNOUNCEMENT)
    assert result.verdict == expected


def test_unrelated_profile_limitations_do_not_make_supported_project_pending():
    output=decision('needs_review','met')
    output['pending_checks']=['税收记录资料未提及','营业执照：仅材料名称或信息需核实，不作为明确硬性冲突依据','无关证件待核实']
    result = validate_decision(BidDecision.model_validate(output), CompanyProfile.model_validate(PROFILE), ANNOUNCEMENT)
    assert result.verdict == "retain" and not result.pending_checks


def test_only_announcement_linked_unknown_mandatory_requirements_need_review():
    output=decision('retain','unknown',quote='')
    output['requirements'] += [dict(output['requirements'][0]),{**output['requirements'][0],'requirement':'本地业绩加分','mandatory':False,'announcement_quote':'评标时具有本地业绩加分。'}]
    output['pending_checks']=['无关食品经营许可证','税收记录资料未提及']
    result=validate_decision(BidDecision.model_validate(output),CompanyProfile.model_validate(PROFILE),ANNOUNCEMENT)
    assert result.verdict=='needs_review' and result.pending_checks==['有效建筑施工许可证：结合证据判断']
    assert len(result.requirements)==3
    output['requirements'][1]['reason']='有效期尚未明确'
    result=validate_decision(BidDecision.model_validate(output),CompanyProfile.model_validate(PROFILE),ANNOUNCEMENT)
    assert result.pending_checks==['有效建筑施工许可证：结合证据判断；有效期尚未明确']
    output['requirements']=output['requirements'][-1:]
    result=validate_decision(BidDecision.model_validate(output),CompanyProfile.model_validate(PROFILE),ANNOUNCEMENT)
    assert result.verdict=='retain' and not result.pending_checks


def test_legacy_checklist_is_read_view_with_general_warnings_once_and_no_mutation(app):
    seed(app,count=2)
    with TestClient(app) as client:
        task_id=create(client).json()['id']
        profile={**PROFILE,'facts':[{**PROFILE['facts'][0],'statement':'营业执照：目录提及。类型：材料名称；主体：其他公司；出现方式：third_party；依据：proof；时间：2024年；行业：食品。'}],
                 'limitations':['税收记录资料未提及']+[f'证件{i}：仅材料名称或信息需核实，不作为明确硬性冲突依据' for i in range(60)]}
        app.state.bid_store.update(task_id,profile=json.dumps(profile),sources=json.dumps([{'source_id':'f1','provenance':[{'ocr':True}]}]),status='completed')
        old=decision('needs_review','met');old['pending_checks']=profile['limitations']
        pending=decision('needs_review','unknown',quote='');pending['pending_checks']=profile['limitations']+['具体核实许可证']
        for i,value in enumerate((old,pending)):app.state.bid_store.item_update(task_id,i,'needs_review',decision=value)
        before=app.state.bid_store.raw(task_id)
        view=client.get('/api/bid-check/tasks/'+task_id).json()
        assert view['has_legacy_decisions'] and view['needs_review']==1
        assert view['items'][0]['decision']['verdict']=='retain' and not view['items'][0]['decision']['pending_checks']
        assert view['items'][1]['decision']['pending_checks']==['有效建筑施工许可证：结合证据判断']
        assert len(view['profile_warnings'])==3
        assert not any('其他公司' in w or '主体' in w for w in view['profile_warnings'])
        assert any('OCR' in w for w in view['profile_warnings'])
        assert not any('不作为明确硬性冲突依据' in w for w in view['profile_warnings'])
        assert app.state.bid_store.raw(task_id)==before
        with app.state.bid_store.connect() as db:
            stored=json.loads(db.execute('SELECT decision FROM items WHERE task_id=? AND position=0',(task_id,)).fetchone()[0])
        assert stored==old
        # Source warnings are derived from the frozen task, never a changed knowledge profile.
        client.put('/api/knowledge/company',json={'name':'另一企业'})
        assert app.state.bid_store.detail(task_id)['profile_warnings']==view['profile_warnings']


def test_new_task_general_warnings_do_not_inflate_project_review_counts(app):
    seed(app,count=2)
    with TestClient(app) as client:
        task_id=create(client).json()['id'];client.post(f'/api/bid-check/tasks/{task_id}/start')
        process(app,task_id,FakeAI([decision('needs_review','met'),decision('needs_review','unknown',quote='')]))
        view=app.state.bid_store.detail(task_id)
        assert view['retained']==2 and view['needs_review']==1 and not view['has_legacy_decisions']
        assert view['profile_warnings'] and not view['items'][0]['decision']['pending_checks']


def test_invented_announcement_and_empty_requirements_fail_not_exclude():
    with pytest.raises(BidCheckError):
        validate_decision(BidDecision.model_validate(decision(reference="编造要求")), CompanyProfile.model_validate(PROFILE), ANNOUNCEMENT)
    with pytest.raises(BidCheckError):
        parse_json_output(json.dumps({**decision(), "requirements": []}), BidDecision)


def test_retry_failed_only_and_stop_resume(app):
    seed(app, count=2)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        ai = FakeAI([BidCheckError("模拟读取失败"), decision("needs_review", "unknown", quote="")])
        process(app, task_id, ai)
        assert client.get(f"/api/bid-check/tasks/{task_id}").json()["failed"] == 1
        assert client.post(f"/api/bid-check/tasks/{task_id}/retry").status_code == 200
        process(app, task_id, ai)
        assert len(ai.calls) == 4  # one profile, two notices, one failed notice retried
        result = client.get(f"/api/bid-check/tasks/{task_id}").json()
        assert result["retained"] == 2 and result["failed"] == 0
        assert client.post(f"/api/bid-check/tasks/{task_id}/retry").status_code == 409


def test_recover_interrupted_running_item_without_repeating_success(app):
    seed(app, count=2)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        app.state.bid_store.update(task_id, profile=json.dumps(PROFILE), status="running")
        app.state.bid_store.item_update(task_id, 0, "retain", decision=decision("retain", "met"))
        app.state.bid_store.item_update(task_id, 1, "running")
        app.state.bid_store.recover()
        assert app.state.bid_store.raw(task_id)["status"] == "queued"
        ai = FakeAI()
        process(app, task_id, ai)
        assert len(ai.calls) == 1 and app.state.bid_store.detail(task_id)["completed"] == 2


def test_stop_during_processing_retains_finished_result_and_continue(app):
    seed(app, count=2)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        class StoppingAI(FakeAI):
            async def generate(self, config, input_text, **kwargs):
                result = await super().generate(config, input_text, **kwargs)
                if "announcement" in json.loads(input_text):
                    app.state.bid_store.update(task_id, stop_requested=1, status="stopping")
                return result
        process(app, task_id, StoppingAI())
        detail = app.state.bid_store.detail(task_id)
        assert detail["status"] == "stopped" and detail["completed"] == 1
        assert client.post(f"/api/bid-check/tasks/{task_id}/start").status_code == 200
        process(app, task_id)
        assert app.state.bid_store.detail(task_id)["completed"] == 2


def test_config_change_blocks_dispatch_and_preserves_files_until_cancelled(app):
    seed(app, count=1)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        app.state.ai_config_store.save(AIConfigUpdate(base_url="https://other.example/v1", model="different", api_key="new"))
        ai = FakeAI()
        process(app, task_id, ai)
        assert not ai.calls
        assert app.state.bid_store.detail(task_id)["status"] == "failed"
        assert (app.state.bid_store.root / task_id).exists()
        assert app.state.bid_store.detail(task_id)["can_retry_preparation"]
        client.post(f"/api/bid-check/tasks/{task_id}/stop")
        assert not (app.state.bid_store.root / task_id).exists()
        assert client.post(f"/api/bid-check/tasks/{task_id}/retry").status_code == 409


def test_transient_profile_error_retries_then_succeeds_without_losing_materials(app):
    seed(app, count=1)
    class FlakyAI(FakeAI):
        attempts = 0
        async def generate(self, config, input_text, **kwargs):
            if "sources" in json.loads(input_text):
                self.attempts += 1
                if self.attempts < 3:
                    raise AIModelError("网关错误（HTTP 502）", status_code=502, category="upstream_server", retryable=True)
            return await super().generate(config, input_text, **kwargs)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        ai = FlakyAI()
        process(app, task_id, ai)
        result = app.state.bid_store.detail(task_id)
        assert ai.attempts == 3 and result["retained"] == 1
        assert not (app.state.bid_store.root / task_id).exists()


def test_exhausted_profile_retry_exposes_safe_diagnostics_and_manual_retry(app, caplog):
    seed(app, count=1)
    class UnavailableAI:
        calls = 0
        async def generate(self, *args, **kwargs):
            self.calls += 1
            raise AIModelError("服务商网关超时（HTTP 504）", status_code=504, category="upstream_server", retryable=True)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        ai = UnavailableAI()
        process(app, task_id, ai)
        result = client.get(f"/api/bid-check/tasks/{task_id}").json()
        assert ai.calls == 3 and result["completed"] == 0
        assert result["can_retry_preparation"] and "已尝试 3 次" in result["error"]
        assert result["diagnostic"]["status_code"] == 504
        assert result["diagnostic"]["phase"] == "company_profile"
        assert result["diagnostic"]["retrying"] is False
        assert TEXT not in caplog.text and "test-secret" not in caplog.text
        assert client.post(f"/api/bid-check/tasks/{task_id}/retry").status_code == 200
        process(app, task_id)
        assert app.state.bid_store.detail(task_id)["retained"] == 1


@pytest.mark.parametrize("category", ["authentication", "compatibility", "rate_limit"])
def test_permanent_profile_errors_do_not_auto_retry(app, category):
    seed(app, count=1)
    class BadAI:
        calls = 0
        async def generate(self, *args, **kwargs):
            self.calls += 1
            raise AIModelError("固定错误", category=category)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        ai = BadAI(); process(app, task_id, ai)
        assert ai.calls == 1 and app.state.bid_store.detail(task_id)["can_retry_preparation"]


def test_stop_during_backoff_does_not_dispatch_another_call(app):
    seed(app, count=1)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        class StopAI:
            calls = 0
            async def generate(self, *args, **kwargs):
                self.calls += 1
                app.state.bid_store.request_stop(task_id)
                raise AIModelError("网关错误", retryable=True)
        ai = StopAI(); process(app, task_id, ai)
        assert ai.calls == 1 and app.state.bid_store.detail(task_id)["status"] == "stopped"
        assert not (app.state.bid_store.root / task_id).exists()


def test_notice_model_outage_pauses_batch_without_repeating_success_on_retry(app):
    seed(app, count=3)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        outage = AIModelError("网关错误（HTTP 502）", status_code=502, retryable=True)
        ai = FakeAI([decision("needs_review", "unknown", quote=""), outage, outage, outage])
        process(app, task_id, ai)
        result = app.state.bid_store.detail(task_id)
        assert result["status"] == "failed" and result["retained"] == 1 and result["failed"] == 1
        assert result["completed"] == 2 and len(ai.calls) == 5
        assert client.post(f"/api/bid-check/tasks/{task_id}/retry").status_code == 200
        process(app, task_id, ai)
        assert len(ai.calls) == 7 and app.state.bid_store.detail(task_id)["retained"] == 3


def test_legacy_failed_task_with_cleaned_uploads_requires_reupload(app):
    seed(app, count=1)
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        app.state.bid_store.update(task_id, status="failed")
        app.state.bid_manager.clean_files(task_id)
        assert not app.state.bid_store.detail(task_id)["can_retry_preparation"]
        assert client.post(f"/api/bid-check/tasks/{task_id}/retry").status_code == 409


def test_deepseek_json_output_runs_profile_and_all_notice_checks(app):
    seed(app, count=2)
    app.state.ai_config_store.save(AIConfigUpdate(base_url="https://api.deepseek.com", model="deepseek-flash", protocol="chat_completions", api_key="deepseek-test-key"))
    calls = []
    def upstream(request):
        payload = json.loads(request.content); calls.append(payload)
        assert payload["response_format"] == {"type": "json_object"}
        assert payload["max_tokens"] == 8192
        assert payload["thinking"] == {"type": "disabled"}
        content = json.loads(payload["messages"][-1]["content"])
        result = PROFILE if "sources" in content else decision("needs_review", "unknown", quote="")
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(result, ensure_ascii=False), "reasoning_content": "not part of final output"}}]})
    with TestClient(app) as client:
        task_id = create(client).json()["id"]
        client.post(f"/api/bid-check/tasks/{task_id}/start")
        process(app, task_id, AIModelClient(transport=httpx.MockTransport(upstream)))
        result = client.get(f"/api/bid-check/tasks/{task_id}").json()
        assert result["status"] == "completed" and result["retained"] == 2
        assert len(calls) == 3 and not result["error"]


@pytest.mark.parametrize("fields", [{"start_date": "2026-10-08", "end_date": "2026-10-07"}, {"region_code": "bad"}, {"keyword": "x" * 101}, {"expected_config": "{}"}])
def test_invalid_creation(app, fields):
    seed(app, count=1)
    with TestClient(app) as client:
        assert create(client, **fields).status_code in {422, 409}
    assert app.state.bid_store.list_tasks() == []
    assert not [path for path in app.state.bid_store.root.iterdir() if path.is_dir()]


def test_no_model_no_files_empty_files_no_results(app):
    with TestClient(app) as client:
        assert create(client).status_code == 422  # no notices
        assert client.post("/api/bid-check/tasks").status_code == 422
        seed(app, count=1)
        assert client.post("/api/bid-check/tasks", files=[("files", ("empty.txt", b"", "text/plain"))]).status_code == 422
        assert client.post("/api/bid-check/tasks", files=[("files", ("bad.exe", b"x"))]).status_code == 422
        app.state.ai_config_store.save(AIConfigUpdate(base_url="https://model.example/v1", model="gpt-test", clear_api_key=True))
        assert create(client).status_code == 422


@pytest.mark.parametrize("status,body,expected", [(403, "blocked", "403"), (200, "captcha", "正文无法识别"), (500, "bad", "无法读取")])
def test_notice_reader_failures(status, body, expected):
    reader = NoticeReader(transport=httpx.MockTransport(lambda request: httpx.Response(status, text=body)))
    with pytest.raises(BidCheckError, match=expected):
        asyncio.run(reader.read("https://www.ccgp.gov.cn/cggg/dfgg/gkzb/202610/a.htm"))


def test_notice_reader_extracts正文_and_rejects_external_urls():
    html = '<html><script>ignore</script><div class="vF_detail_content">' + ANNOUNCEMENT * 3 + '</div></html>'
    reader = NoticeReader(transport=httpx.MockTransport(lambda request: httpx.Response(200, text=html)))
    text = asyncio.run(reader.read("https://www.ccgp.gov.cn/cggg/dfgg/gkzb/202610/a.htm"))
    assert "申请人" in text and "ignore" not in text
    with pytest.raises(BidCheckError):
        asyncio.run(reader.read("http://127.0.0.1/secrets"))


def test_document_text_pdf_word_spreadsheet_and_images(tmp_path):
    (tmp_path / "1.txt").write_text(TEXT, encoding="utf-8")
    doc = Document(); doc.add_paragraph(TEXT); doc.add_table(rows=1, cols=1).cell(0, 0).text = "资质说明"
    doc.save(tmp_path / "2.docx")
    book = Workbook(); book.active["A1"] = "营业范围"; book.active["B1"] = "设备"; book.save(tmp_path / "3.xlsx")
    Image.new("RGB", (100, 80), "white").save(tmp_path / "4.png")
    pdf = pymupdf.open(); page = pdf.new_page(); page.insert_text((40, 40), "Company evidence in a readable PDF document " * 2); pdf.save(tmp_path / "5.pdf"); pdf.close()
    scan = pymupdf.open(); scan.new_page(); scan.save(tmp_path / "6.pdf"); scan.close()
    sources, images = extract_documents([{"name": name, "stored_name": name} for name in ["1.txt", "2.docx", "3.xlsx", "4.png", "5.pdf", "6.pdf"]], tmp_path)
    assert any("资质说明" in item["text"] for item in sources)
    assert any("B1=设备" in item["text"] for item in sources)
    assert len(images) == 2
    assert len({item["source_id"] for item in sources}) == len(sources)


def test_corrupt_documents_not_silently_ignored(tmp_path):
    (tmp_path / "bad.pdf").write_bytes(b"not-pdf")
    with pytest.raises(BidCheckError):
        extract_documents([{"name": "bad.pdf", "stored_name": "bad.pdf"}], tmp_path)


@pytest.mark.parametrize("protocol", ["responses", "chat_completions"])
def test_multimodal_generation_schema_and_compatible_fallback(app, protocol):
    app.state.ai_config_store.save(AIConfigUpdate(base_url="https://model.example/v1", model="gpt-test", protocol=protocol))
    calls = []
    def upstream(request):
        body = json.loads(request.content); calls.append(body)
        if len(calls) == 1:
            return httpx.Response(400, json={"error": {"message": "json_schema not supported"}})
        content = body["input"][0]["content"] if protocol == "responses" else body["messages"][-1]["content"]
        assert any(item["type"] in {"input_image", "image_url"} for item in content)
        assert "response_format" not in body and "text" not in body
        return httpx.Response(200, json={"output": [{"type": "message", "content": [{"type": "output_text", "text": "{}"}]}]} if protocol == "responses" else {"choices": [{"message": {"content": "{}"}}]})
    client = AIModelClient(transport=httpx.MockTransport(upstream))
    asyncio.run(client.generate(app.state.ai_config_store.get(), "公司资料", images=[{"source_id": "image1", "data_url": "data:image/png;base64,test"}], schema=CompanyProfile.model_json_schema(), timeout=120))
    assert len(calls) == 2
