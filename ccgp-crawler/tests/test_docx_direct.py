"""Structural coverage and cancellation without live model or Office dependencies."""
import asyncio
import io
import json
import signal
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from ccgp_crawler import bid_documents as office
from ccgp_crawler.bid_models import BidCheckError
from ccgp_crawler.bid_service import TaskInterrupted
from ccgp_crawler.docx_reader import DocxReader, W, R, A, MC
from ccgp_crawler.knowledge_documents import index_document, image_bytes
from test_knowledge import app, CONFIG, KnowledgeAI, entry
from ccgp_crawler.knowledge_models import ExtractionOutput


def package(path, body, parts=None, rels=''):
    namespaces=f'xmlns:w="{W}" xmlns:r="{R}" xmlns:a="{A}" xmlns:mc="{MC}" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"'
    with zipfile.ZipFile(path,'w') as z:
        z.writestr('word/document.xml',f'<w:document {namespaces}><w:body>{body}</w:body></w:document>')
        z.writestr('word/_rels/document.xml.rels',f'<Relationships>{rels}</Relationships>')
        for key,value in (parts or {}).items():z.writestr(key,value)
    return path


def picture():
    stream=io.BytesIO();Image.new('RGB',(50,50),'white').save(stream,'PNG');return stream.getvalue()


def test_direct_structural_coverage_and_compatibility(tmp_path,monkeypatch):
    p=package(tmp_path/'file.docx','''<w:p><w:r><w:t>标题</w:t></w:r></w:p>
    <w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>合并表格</w:t></w:r></w:p>
    <w:tbl><w:tr><w:tc><w:p><w:r><w:t>嵌套表格</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:tc></w:tr></w:tbl>
    <w:p><w:r><w:t>图片附近</w:t><mc:AlternateContent><mc:Choice><a:blip r:embed="img"/></mc:Choice><mc:Fallback><v:imagedata r:id="img"/></mc:Fallback></mc:AlternateContent></w:r></w:p>
    <w:p><w:r><v:imagedata r:id="img"/><w:txbxContent><w:p><w:r><w:t>文本框文字</w:t></w:r></w:p></w:txbxContent>
    <w:footnoteReference w:id="1"/><w:endnoteReference w:id="2"/><dgm:relIds r:dm="diagram"/></w:r></w:p>
    <w:sectPr><w:headerReference r:id="header"/><w:footerReference r:id="footer"/></w:sectPr>''',{
        'word/media/img.dat':picture(),'word/media/unused.png':picture(),
        'word/header1.xml':f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>页眉企业</w:t></w:r></w:p></w:hdr>',
        'word/footer1.xml':f'<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>页脚日期</w:t></w:r></w:p></w:ftr>',
        'word/footnotes.xml':f'<w:footnotes xmlns:w="{W}"><w:footnote w:id="1"><w:p><w:r><w:t>脚注</w:t></w:r></w:p></w:footnote></w:footnotes>',
        'word/endnotes.xml':f'<w:endnotes xmlns:w="{W}"><w:endnote w:id="2"><w:p><w:r><w:t>尾注</w:t></w:r></w:p></w:endnote></w:endnotes>',
        'word/diagrams/data1.xml':f'<dgm:data xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="{A}" xmlns:r="{R}"><dgm:pt><a:t>SmartArt能力</a:t><a:blip r:embed="di"/></dgm:pt></dgm:data>',
        'word/diagrams/_rels/data1.xml.rels':'<Relationships><Relationship Id="di" Target="../media/img.dat"/></Relationships>'
    }, ''.join(f'<Relationship Id="{rid}" Target="{target}"/>' for rid,target in [('img','media/img.dat'),('header','header1.xml'),('footer','footer1.xml'),('diagram','diagrams/data1.xml')]))
    monkeypatch.setattr(office,'soffice_path',lambda:None)
    units,stats=DocxReader(p).read()
    text='\n'.join(u['text'] for u in units)
    for expected in ['标题','合并表格','嵌套表格','文本框文字','页眉企业','页脚日期','脚注','尾注','SmartArt能力']:assert expected in text
    assert text.count('文本框文字')==1
    assert stats['tables']==2 and stats['images']==3 and stats['unique_images']==1 and stats['unreferenced_media']==1
    assert len({u['location'] for u in units if u['visual']})==3
    assert any('合并2列' in u['location'] for u in units)
    pages,batches=index_document(p,p.name,tmp_path/'work');assert pages==len(units)
    assert all(len(b)<=6 and sum(len(u['text']) for u in b)<=20000 for b in batches)


def test_short_paragraph_merging_large_paragraph_and_precise_sources(tmp_path):
    doc=Document()
    for _ in range(4000):doc.add_paragraph('短段落')
    doc.add_paragraph('长'*9000);p=tmp_path/'long.docx';doc.save(p)
    units,stats=DocxReader(p).read()
    assert stats['paragraphs']==4001 and len(units)<20
    assert sum(len(segment['text']) for u in units for segment in u['segments'])==21000
    assert all(len(u['text'])<=4000 for u in units)


def test_images_large_limit_corruption_and_no_partial_batches(tmp_path):
    large=tmp_path/'large.png';Image.new('L',(8000,6000)).save(large)
    data=image_bytes({'kind':'image','path':str(large)})
    with Image.open(io.BytesIO(data)) as im:assert max(im.size)<=2400
    oversized=tmp_path/'oversized.png';Image.new('1',(10001,10000)).save(oversized)
    with pytest.raises(BidCheckError,match='1 亿'):image_bytes({'kind':'image','path':str(oversized)})
    broken=package(tmp_path/'broken.docx','<w:p><a:blip r:embed="bad"/></w:p>',{'word/media/bad.png':b'bad'},'<Relationship Id="bad" Target="media/bad.png"/>')
    progress=[]
    with pytest.raises(BidCheckError,match='正文第1段.*图片损坏'):index_document(broken,'broken.docx',tmp_path,progress=lambda **kw:progress.append(kw))
    assert progress[-1]['unreadable'] and all(p['phase']!='indexing_complete' for p in progress)


def test_direct_read_stop(tmp_path):
    p=package(tmp_path/'file.docx','<w:p><w:r><w:t>内容</w:t></w:r></w:p>')
    def check():raise TaskInterrupted()
    with pytest.raises(TaskInterrupted):index_document(p,p.name,tmp_path,check=check)


def test_compatibility_uses_readable_fallback_once(tmp_path):
    path=package(tmp_path/'fallback.docx','''<w:p><mc:AlternateContent>
    <mc:Choice><a:blip r:embed="modern"/></mc:Choice>
    <mc:Fallback><v:imagedata r:id="fallback"/></mc:Fallback>
    </mc:AlternateContent></w:p>''',{'word/media/modern.svg':b'<svg/>','word/media/fallback.dat':picture()},
    '<Relationship Id="modern" Target="media/modern.svg"/><Relationship Id="fallback" Target="media/fallback.dat"/>')
    units,stats=DocxReader(path).read()
    assert stats['images']==1 and units[0]['visual']['media']=='word/media/fallback.dat'


@pytest.mark.parametrize('timeout', ['1200','60','3600','invalid'])
def test_conversion_timeout_and_cleanup(tmp_path,monkeypatch,timeout):
    monkeypatch.setattr(office,'soffice_path',lambda:'office')
    monkeypatch.setenv('CCGP_OFFICE_CONVERT_TIMEOUT_SECONDS',timeout)
    values=iter([0,4000]);monkeypatch.setattr(office.time,'monotonic',lambda:next(values))
    killed=[];monkeypatch.setattr(office.os,'killpg',lambda pid,sig:killed.append(sig))
    class Process:
        pid=123
        def communicate(self,timeout=None):return b'',b''
    monkeypatch.setattr(office.subprocess,'Popen',lambda *a,**kw:Process())
    expected=timeout if timeout.isdigit() else '1200'
    with pytest.raises(BidCheckError,match=expected):office._convert(tmp_path/'a.doc','pdf',tmp_path)
    assert killed==[signal.SIGTERM] and not list(tmp_path.glob('office-convert-*'))


def test_conversion_stop_exit_failure_and_isolated_retries(tmp_path,monkeypatch):
    monkeypatch.setattr(office,'soffice_path',lambda:'office');dirs=[];killed=[]
    monkeypatch.setattr(office.os,'killpg',lambda *args:killed.append(args))
    class Process:
        pid=123;returncode=0
        def communicate(self,timeout=None):return b'',b''
    def launch(args,**kwargs):
        assert kwargs['start_new_session']
        output=Path(args[args.index('--outdir')+1]);dirs.append(output);(output/'a.pdf').write_bytes(b'%PDF')
        return Process()
    monkeypatch.setattr(office.subprocess,'Popen',launch)
    one=office._convert(tmp_path/'a.doc','pdf',tmp_path);two=office._convert(tmp_path/'a.doc','pdf',tmp_path)
    assert one.parent!=two.parent
    calls=0
    def stop():
        nonlocal calls
        calls+=1
        if calls>1:raise TaskInterrupted()
    with pytest.raises(TaskInterrupted):office._convert(tmp_path/'a.doc','pdf',tmp_path,check=stop)
    assert killed and not dirs[-1].exists()
    Process.returncode=2
    with pytest.raises(BidCheckError,match='退出失败'):office._convert(tmp_path/'a.doc','pdf',tmp_path)
    assert not dirs[-1].exists()


def test_zero_batch_retry_does_not_require_model_and_waits_for_start(tmp_path,app):
    test_app=app;test_app.state.ai_client=KnowledgeAI()
    store=test_app.state.knowledge_store;store.save_company('测试公司')
    path=tmp_path/'original.docx';doc=Document();doc.add_paragraph('证照主体');doc.save(path)
    doc_id,_=store.add_document('doc','原件.docx','hash',path.stat().st_size,path)
    run=store.document(doc_id)['run'];store.update_run(run['id'],status='failed',error='旧错误')
    store.action(doc_id,'retry')
    asyncio.run(test_app.state.knowledge_manager.process_run(run['id']))
    reread=store.document(doc_id)['run'];assert reread['status']=='uploaded' and reread['total_batches']==1
    assert test_app.state.ai_client.calls==[]
    assert reread['diagnostic']['mode']=='docx_direct'


def test_quote_source_maps_to_actual_paragraph(tmp_path,app):
    test_app=app;store=test_app.state.knowledge_store;store.save_company('测试公司')
    path=tmp_path/'original.txt';path.write_text('其他\n证照主体')
    doc_id,_=store.add_document('doc','原件.txt','hash',path.stat().st_size,path)
    run=store.document(doc_id)['run'];unit={'source_id':'S1','text':'其他\n证照主体','location':'第1段至第2段','visual':None,'segments':[{'text':'其他','location':'第1段'},{'text':'证照主体','location':'第2段'}]}
    store.set_batches(run['id'],1,[[unit]])
    output=ExtractionOutput.model_validate({'entries':[entry(quote='证照主体')],'limitations':[]})
    store.finish_batch(run['id'],0,output,[unit])
    assert store.entries()['items'][0]['source']['location']=='第2段'


def test_legacy_doc_xls_and_unsupported_docx_use_fallback(tmp_path,monkeypatch):
    import pymupdf
    from openpyxl import Workbook
    from ccgp_crawler import knowledge_documents as documents
    calls=[]
    def convert(source,target,work,**kwargs):
        calls.append((source.suffix,target))
        kwargs['check']()
        output=work/f'converted.{target}'
        if target=='pdf':
            with pymupdf.open() as pdf:
                page=pdf.new_page();page.insert_text((70,70),'Converted document text '*10);pdf.save(output)
        else:
            book=Workbook();book.active['A1']='企业名称';book.save(output)
        return output
    monkeypatch.setattr(documents,'_convert',convert)
    old=tmp_path/'old.doc';old.write_bytes(b'legacy')
    assert index_document(old,old.name,tmp_path)[1]
    excel=tmp_path/'old.xls';excel.write_bytes(b'legacy')
    assert index_document(excel,excel.name,tmp_path)[1]
    graph=package(tmp_path/'graph.docx','<w:p><w:r><w:object/></w:r></w:p>')
    reports=[]
    assert index_document(graph,graph.name,tmp_path,progress=lambda **kw:reports.append(kw))[1]
    assert calls==[('.doc','pdf'),('.xls','xlsx'),('.docx','pdf')]
    assert reports[-1]['mode']=='office_convert' and reports[-1]['unreadable']==[]


def test_api_reread_without_key_and_unknown_document_404(app,tmp_path):
    from fastapi.testclient import TestClient
    from ccgp_crawler.ai_config import AIConfigUpdate
    store=app.state.knowledge_store;store.save_company('测试公司')
    path=tmp_path/'original.txt';path.write_text('企业材料')
    doc_id,_=store.add_document('doc','原件.txt','hash',path.stat().st_size,path)
    run=store.document(doc_id)['run'];store.update_run(run['id'],status='failed')
    app.state.ai_config_store.save(AIConfigUpdate(**CONFIG,clear_api_key=True))
    with TestClient(app) as client:
        assert client.post(f'/api/knowledge/documents/{doc_id}/action',json={'action':'retry'}).status_code==200
        assert client.post('/api/knowledge/documents/missing/action',json={'action':'retry'}).status_code==404
    asyncio.run(app.state.knowledge_manager.process_run(run['id']))
    assert store.document(doc_id)['run']['status']=='uploaded'
    with TestClient(app) as client:
        assert client.post(f'/api/knowledge/documents/{doc_id}/action',json={'action':'start','expected_config':CONFIG}).status_code==422
