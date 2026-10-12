"""Knowledge coverage, evidence review and immutable bid snapshots (no live AI)."""
import asyncio
import io
import json
import stat
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from ccgp_crawler.bid_models import BidCheckError, BidDecision, CompanyProfile, validate_decision
from ccgp_crawler.knowledge_documents import index_document, batch_images
from ccgp_crawler.knowledge_models import ExtractionOutput, validate_extraction
from ccgp_crawler.knowledge_store import KnowledgeStorageError
from ccgp_crawler.ai_client import AIModelError
from test_bid_check import app, seed, FakeReader, ANNOUNCEMENT

CONFIG = {'base_url':'https://model.example/v1','model':'gpt-test','protocol':'responses'}
TEXT = '测试公司具有独立承担民事责任的能力。'


def entry(**changes):
    return {'kind':'fact','category':'资格声明','name':'民事责任能力','value':'具有独立承担民事责任的能力','subject':'测试公司','occurrence':'actual','evidence_kind':'declaration','source_id':'S1','quote':TEXT,'applicable_date':'','industry':'','notes':'','needs_review':False,**changes}


class KnowledgeAI:
    def __init__(self, stop=None):self.calls=[];self.stop=stop
    async def generate(self, config, input_text, **kwargs):
        data=json.loads(input_text);self.calls.append((data,kwargs))
        if self.stop:self.stop()
        if 'sources' in data:
            source=data['sources'][0]
            return json.dumps({'entries':[entry(source_id=source['source_id'],quote=source['text'][:100] or '扫描证件')], 'limitations':[]},ensure_ascii=False)
        if 'clause_index' in data:
            return json.dumps({'requirements':[], 'groups':[]})
        return json.dumps({'verdict':'needs_review','summary':'待核实','pending_checks':[],'requirements':[{'requirement':'施工许可证','mandatory':True,'announcement_quote':'供应商必须具备有效的建筑施工许可证。','assessment':'unknown','company_fact_ids':[],'company_quote':'','reason':'资料缺失不是明确冲突'}]},ensure_ascii=False)


def upload(client, text=TEXT, name='历史投标.txt'):
    return client.post('/api/knowledge/documents',files=[('files',(name,text.encode(),'text/plain'))])


def setup_document(app, client, text=TEXT):
    client.put('/api/knowledge/company',json={'name':'测试公司'})
    response=upload(client,text);assert response.status_code==200,response.text
    doc=response.json()['items'][0]['document']
    asyncio.run(app.state.knowledge_manager.process_run(doc['run']['id']))
    return app.state.knowledge_store.document(doc['id'])


def extract(app,client,doc,ai=None):
    manager=app.state.knowledge_manager;manager.interval=0;manager.retry_delays=(0,0)
    app.state.ai_client=ai or KnowledgeAI()
    response=client.post(f"/api/knowledge/documents/{doc['id']}/action",json={'action':'start','expected_config':CONFIG})
    assert response.status_code==200,response.text
    asyncio.run(manager.process_run(doc['run']['id']))
    return client.get('/api/knowledge/entries',params={'document_id':doc['id']}).json()['items']


def confirm(client,items):
    r=client.post('/api/knowledge/review',json={'items':[{'id':i['id'],'status':'accepted'} for i in items]})
    assert r.status_code==200,r.text
    return r.json()['revision_id']


def test_upload_name_dedup_new_version_permissions_and_backup(app):
    with TestClient(app) as client:
        assert upload(client).status_code==422
        doc=setup_document(app,client)
        path,_=app.state.knowledge_store.original_path(doc['id'])
        assert path.read_text()==TEXT and stat.S_IMODE(path.stat().st_mode)==0o600
        assert stat.S_IMODE(app.state.knowledge_store.root.stat().st_mode)==0o700
        result=upload(client).json()['items'][0]
        assert result['duplicate'] and result['document']['id']==doc['id']
        assert len(app.state.knowledge_store.documents())==1
        assert not hasattr(app.state.ai_client,'calls')
        assert upload(client,TEXT+'修改后的内容').json()['items'][0]['document']['id']!=doc['id']
        assert client.get(f"/api/knowledge/documents/{doc['id']}/original").content==TEXT.encode()
        with zipfile.ZipFile(io.BytesIO(client.get('/api/knowledge/backup').content)) as backup:
            assert 'knowledge.sqlite3' in backup.namelist()
            assert len([n for n in backup.namelist() if '/original.' in n])==2


def test_complete_extract_review_audit_and_revision_immutable(app):
    with TestClient(app) as client:
        doc=setup_document(app,client);items=extract(app,client,doc)
        assert app.state.knowledge_store.document(doc['id'])['run']['status']=='awaiting_review'
        assert len(items)==1 and items[0]['status']=='pending'
        empty=client.get('/api/knowledge/company').json()['revision']['id']
        assert client.get('/api/knowledge/revisions/'+empty).status_code==422
        revision=confirm(client,items)
        original=app.state.knowledge_store.revision(revision)
        assert original['fact_count']==1
        client.post('/api/knowledge/review',json={'items':[{'id':items[0]['id'],'status':'accepted','changes':{'value':'人工修改能力','needs_review':True}}]})
        changed=client.get('/api/knowledge/entries').json()['items'][0]
        assert changed['original']['value']!=changed['data']['value']
        assert app.state.knowledge_store.revision(revision)==original
        with app.state.knowledge_store.connect() as db:assert db.execute('SELECT count(*) FROM review_events').fetchone()[0]==2
        client.put(f"/api/knowledge/documents/{doc['id']}/active",json={'active':False})
        assert client.get('/api/knowledge/company').json()['revision']['fact_count']==0
        assert app.state.knowledge_store.revision(revision)==original and app.state.knowledge_store.original_path(doc['id'])[0].is_file()


@pytest.mark.parametrize('occurrence,evidence_kind,subject',[('mention','proof','测试公司'),('template','declaration','测试公司'),('actual','proof','制造商'),('actual','inference','测试公司'),('actual','proof','')])
def test_unverified_classifications_not_company_proof(occurrence,evidence_kind,subject):
    result=validate_extraction(ExtractionOutput.model_validate({'entries':[entry(occurrence=occurrence,evidence_kind=evidence_kind,subject=subject)],'limitations':[]}),[{'source_id':'S1','text':TEXT,'visual':None}], '测试公司')
    assert result.entries[0].needs_review
    if subject=='制造商':assert result.entries[0].occurrence=='actual' and result.entries[0].subject_role=='unknown'


def test_invalid_source_and_quote_rejected():
    for changes in [{'source_id':'invented'},{'quote':'编造内容'}]:
        with pytest.raises(BidCheckError):
            validate_extraction(ExtractionOutput.model_validate({'entries':[entry(**changes)],'limitations':[]}),[{'source_id':'S1','text':TEXT,'visual':None}], '测试公司')


def test_sme_subject_industry_date_material_separate(app):
    with TestClient(app) as client:
        doc=setup_document(app,client)
        batch=app.state.knowledge_store.next_batch(doc['run']['id'])
        output=ExtractionOutput.model_validate({'entries':[entry(kind='material',category='企业规模',name='中小企业声明函'),entry(category='企业规模',name='企业规模',value='小型企业',applicable_date='2020-10-01',industry='制造业')],'limitations':[]})
        app.state.knowledge_store.finish_batch(doc['run']['id'],0,output,batch['units'])
        app.state.knowledge_store.update_run(doc['run']['id'],status='awaiting_review')
        items=app.state.knowledge_store.entries()['items'];revision=confirm(client,items)
        snapshot=app.state.knowledge_store.revision(revision)
        assert snapshot['material_count']==snapshot['fact_count']==1
        assert snapshot['records'][1]['data']['industry']=='制造业'
        assert snapshot['records'][1]['data']['applicable_date']=='2020-10-01'
        assert snapshot['sources'][0]['unsafe_for_conflict']


def test_stop_resume_and_recovery_only_pending_batches(app):
    with TestClient(app) as client:
        doc=setup_document(app,client,TEXT*2000)
        assert doc['run']['total_batches']>1
        manager=app.state.knowledge_manager
        ai=KnowledgeAI(stop=lambda: app.state.knowledge_store.action(doc['id'],'stop'))
        extract(app,client,doc,ai)
        assert len(ai.calls)==1
        assert app.state.knowledge_store.document(doc['id'])['run']['status']=='stopped'
        ai.stop=None
        extract(app,client,doc,ai)
        assert len(ai.calls)==doc['run']['total_batches']
        assert app.state.knowledge_store.document(doc['id'])['run']['status']=='awaiting_review'
        # Interrupted batches return to pending while successful batches remain done.
        app.state.knowledge_store.update_run(doc['run']['id'],status='running')
        app.state.knowledge_store.batch_state(doc['run']['id'],1,'running')
        app.state.knowledge_store.recover()
        assert app.state.knowledge_store.raw_run(doc['run']['id'])['status']=='queued'
        assert app.state.knowledge_store.next_batch(doc['run']['id'])['position']==1
        pending_entry=next(i for i in app.state.knowledge_store.entries()['items'] if i['position']==1)
        assert not pending_entry['reviewable']
        assert client.post('/api/knowledge/review',json={'items':[{'id':pending_entry['id'],'status':'accepted'}]}).status_code==422


def test_transient_retry_failed_batch_and_reextract_old_revision(app):
    with TestClient(app) as client:
        doc=setup_document(app,client)
        class Flaky(KnowledgeAI):
            attempts=0
            async def generate(self,*args,**kwargs):
                self.attempts+=1
                if self.attempts<3:raise AIModelError('网关错误',retryable=True)
                return await super().generate(*args,**kwargs)
        ai=Flaky();items=extract(app,client,doc,ai);assert ai.attempts==3
        revision=confirm(client,items)
        result=client.post(f"/api/knowledge/documents/{doc['id']}/action",json={'action':'reextract'})
        assert result.status_code==200 and result.json()['run']['generation']==2
        assert client.get('/api/knowledge/company').json()['revision']['id']==revision
        newdoc=result.json();asyncio.run(app.state.knowledge_manager.process_run(newdoc['run']['id']))
        extract(app,client,newdoc)
        # Unconfirmed new extraction does not replace the confirmed original.
        assert client.get('/api/knowledge/company').json()['revision']['id']==revision
        client.put('/api/knowledge/company',json={'name':'测试公司'})
        assert client.get('/api/knowledge/company').json()['revision']['fact_count']==1


def test_partial_confirmation_and_edit_survive_retry_without_reprocessing(app):
    with TestClient(app) as client:
        doc=setup_document(app,client,TEXT*2000)
        class Outage(KnowledgeAI):
            count=0
            async def generate(self,*args,**kwargs):
                self.count+=1
                if self.count>1:raise AIModelError('超时',retryable=True)
                return await super().generate(*args,**kwargs)
        ai=Outage();items=extract(app,client,doc,ai)
        assert len(items)==1 and ai.count==4
        store=app.state.knowledge_store
        assert items[0]['reviewable'] and items[0]['batch_status']=='done'
        result=client.post('/api/knowledge/review',json={'items':[{'id':items[0]['id'],'status':'accepted','changes':{'value':'人工核对能力','needs_review':True,'notes':'需核对有效期'}}]})
        assert result.status_code==200,result.text
        revision=result.json()['revision_id'];original=store.revision(revision)
        assert original['fact_count']==1 and original['records'][0]['data']['needs_review']
        assert original['unsafe_fact_ids'] and any('尚未完整拆解' in v for v in original['profile']['limitations'])
        assert store.document(doc['id'])['run']['status']=='failed'
        response=client.post(f"/api/knowledge/documents/{doc['id']}/action",json={'action':'retry','expected_config':CONFIG})
        assert response.status_code==200
        # A new snapshot while queued must still contain the adopted successful batch.
        client.put('/api/knowledge/company',json={'name':'测试公司'})
        assert client.get('/api/knowledge/company').json()['revision']['fact_count']==1
        good=KnowledgeAI();app.state.ai_client=good
        asyncio.run(app.state.knowledge_manager.process_run(doc['run']['id']))
        assert len(good.calls)==doc['run']['total_batches']-1
        retained=next(i for i in store.entries()['items'] if i['id']==items[0]['id'])
        assert retained['status']=='accepted' and retained['data']['value']=='人工核对能力'
        assert retained['data']['needs_review'] and retained['original']['value']!=retained['data']['value']
        assert store.revision(revision)==original
        current=client.get('/api/knowledge/company').json()['revision']
        assert current['fact_count']==1 and not any('尚未完整拆解' in v for v in current['profile']['limitations'])
        assert store.document(doc['id'])['run']['status']=='awaiting_review'
        # Even an accidental duplicate completion cannot erase an accepted entry.
        store.finish_batch(doc['run']['id'],0,ExtractionOutput(entries=[],limitations=[]),[])
        assert next(i for i in store.entries()['items'] if i['id']==items[0]['id'])['status']=='accepted'


@pytest.mark.parametrize('status',['running','queued','stopped','stopping','failed'])
def test_review_successful_batch_preserves_incomplete_run_and_rejects_invalid_batch(app,status):
    with TestClient(app) as client:
        doc=setup_document(app,client,TEXT*2000);store=app.state.knowledge_store;run_id=doc['run']['id']
        for _ in range(2):
            batch=store.next_batch(run_id)
            output=ExtractionOutput.model_validate({'entries':[entry(source_id=batch['units'][0]['source_id'])],'limitations':[]})
            store.finish_batch(run_id,batch['position'],output,batch['units'])
        store.batch_state(run_id,1,'failed','引用不符')
        store.update_run(run_id,status=status,error='原错误')
        good,bad=store.entries()['items']
        assert good['reviewable'] and not bad['reviewable']
        # Mixed batch review is atomic: no candidate is adopted if another is invalid.
        rejected=client.post('/api/knowledge/review',json={'items':[{'id':good['id'],'status':'accepted'},{'id':bad['id'],'status':'accepted'}]})
        assert rejected.status_code==422 and '所在批次' in rejected.json()['detail']
        assert store.entries()['items'][0]['status']=='pending'
        for action in ('accepted','rejected','pending','accepted'):
            response=client.post('/api/knowledge/review',json={'items':[{'id':good['id'],'status':action}]})
            assert response.status_code==200,response.text
            run=store.document(doc['id'])['run']
            assert run['status']==status and run['error']=='原错误'
        with store.connect() as db:db.execute('DELETE FROM batches WHERE run_id=? AND position=1',(run_id,))
        assert not store.entries()['items'][1]['reviewable']
        assert client.post('/api/knowledge/review',json={'items':[{'id':bad['id'],'status':'rejected'}]}).status_code==422


def test_finishing_after_early_review_updates_coverage_without_overwriting_revision(app):
    with TestClient(app) as client:
        doc=setup_document(app,client,TEXT*2000);store=app.state.knowledge_store;run_id=doc['run']['id']
        batch=store.next_batch(run_id)
        store.finish_batch(run_id,0,ExtractionOutput.model_validate({'entries':[entry()],'limitations':[]}),batch['units'])
        store.update_run(run_id,status='running')
        revision=confirm(client,store.entries()['items']);original=store.revision(revision)
        while batch:=store.next_batch(run_id):
            store.finish_batch(run_id,batch['position'],ExtractionOutput(entries=[],limitations=[]),batch['units'])
        store.finish_run(run_id)
        assert store.document(doc['id'])['run']['status']=='reviewed'
        assert not any('尚未完整拆解' in v for v in store.company()['revision']['profile']['limitations'])
        assert store.revision(revision)==original


def test_duplicate_facts_merge_and_conflicts_remain_parallel(app):
    with TestClient(app) as client:
        for text in [TEXT,TEXT+'另一出处',TEXT+'冲突出处']:
            doc=setup_document(app,client,text);items=extract(app,client,doc)
            changes={'value':'不具有该能力'} if text.endswith('冲突出处') else {}
            client.post('/api/knowledge/review',json={'items':[{'id':items[0]['id'],'status':'accepted','changes':changes}]})
        snapshot=client.get('/api/knowledge/company').json()['revision']
        assert len(snapshot['records'])==2
        assert sorted(len(r['sources']) for r in snapshot['records'])==[1,2]
        assert all(r['conflict'] for r in snapshot['records'])


def test_two_bid_tasks_reuse_without_profile_call_or_mutation(app):
    seed(app,count=2)
    with TestClient(app) as client:
        doc=setup_document(app,client);items=extract(app,client,doc);revision=confirm(client,items)
        ai=KnowledgeAI();app.state.ai_client=ai;app.state.bid_manager.reader=FakeReader()
        for _ in range(2):
            created=client.post('/api/bid-check/tasks',data={'knowledge_revision_id':revision,'expected_config':json.dumps(CONFIG)}).json()
            assert created['estimated_calls']==created['total']==2
            client.post('/api/bid-check/tasks/'+created['id']+'/start')
            asyncio.run(app.state.bid_manager.process_task(created['id']))
            task=app.state.bid_store.raw(created['id']);assert task['knowledge_revision_id']==revision and task['profile']
            assert app.state.bid_store.detail(created['id'])['retained']==2 and task['evidence_index']
        assert len(ai.calls)==4 and all('sources' not in call[0] for call in ai.calls)
        before=app.state.bid_store.raw(created['id'])['profile']
        client.put(f"/api/knowledge/documents/{doc['id']}/active",json={'active':False})
        assert app.state.bid_store.raw(created['id'])['profile']==before
        assert client.post('/api/bid-check/tasks',data={'knowledge_revision_id':revision},files=[('files',('x.txt',b'x'))]).status_code==422


def test_formats_and_hundreds_of_pages_complete_coverage(tmp_path):
    work=tmp_path/'work';work.mkdir()
    text=TEXT*3000;(tmp_path/'a.txt').write_text(text)
    _,batches=index_document(tmp_path/'a.txt','a.txt',work)
    assert ''.join(u['text'] for b in batches for u in b)==text
    assert all(len(b)<=6 and sum(len(u['text']) for u in b)<=20000 for b in batches)
    doc=Document();doc.add_paragraph(TEXT);doc.add_table(rows=1,cols=1).cell(0,0).text=TEXT;doc.save(tmp_path/'b.docx')
    pages,batches=index_document(tmp_path/'b.docx','b.docx',work);assert pages>=1 and any(u['text'].strip() for b in batches for u in b)
    Image.new('RGB',(80,80),'white').save(tmp_path/'c.png')
    book=Workbook();book.active['A1']=TEXT;book.active.add_image(__import__('openpyxl').drawing.image.Image(tmp_path/'c.png'));book.save(tmp_path/'d.xlsx')
    _,batches=index_document(tmp_path/'d.xlsx','d.xlsx',work);assert any('A1=' in u['text'] for b in batches for u in b) and sum(len(batch_images(b)) for b in batches)==1
    pdf=pymupdf.open()
    for n in range(240):
        page=pdf.new_page();page.insert_text((40,40),('Readable qualification evidence in document. '*3) if n%2==0 else '')
    pdf.save(tmp_path/'e.pdf');pdf.close()
    pages,batches=index_document(tmp_path/'e.pdf','e.pdf',work)
    assert pages==240 and sum(len(b) for b in batches)==240
    assert sum(bool(u['visual']) for b in batches for u in b)==120
    assert len(batch_images(batches[0]))==3
    assert not list(work.rglob('*.png')) and not list(work.rglob('*.jpg'))


def test_oversized_page_count_and_image_volume_split(tmp_path,monkeypatch):
    work=tmp_path/'work';work.mkdir();pdf=pymupdf.open()
    for _ in range(3001):pdf.new_page(width=100,height=100)
    pdf.save(tmp_path/'large.pdf');pdf.close()
    with pytest.raises(BidCheckError,match='3,000'):index_document(tmp_path/'large.pdf','large.pdf',work)
    from ccgp_crawler import knowledge_documents as module
    pdf=pymupdf.open()
    for _ in range(6):pdf.new_page(width=100,height=100)
    pdf.save(tmp_path/'small.pdf');pdf.close()
    monkeypatch.setattr(module,'image_bytes',lambda visual:b'x'*(4*1024*1024))
    _,batches=index_document(tmp_path/'small.pdf','small.pdf',work)
    assert len(batches)==3 and [len(b) for b in batches]==[2,2,2]


def test_store_error_safe_503_and_search(app,monkeypatch):
    with TestClient(app) as client:
        doc=setup_document(app,client);extract(app,client,doc)
        assert client.get('/api/knowledge/entries',params={'q':'民事责任'}).json()['total']==1
        assert client.get('/api/knowledge/entries',params={'category':'企业规模'}).json()['total']==0
        monkeypatch.setattr(app.state.knowledge_store,'company',lambda: (_ for _ in ()).throw(KnowledgeStorageError('知识库存储不可用')))
        result=client.get('/api/knowledge/company');assert result.status_code==503
        assert 'test-secret' not in result.text


def test_withdraw_new_adopted_run_never_resurrects_old_entries(app):
    with TestClient(app) as client:
        doc=setup_document(app,client);items=extract(app,client,doc);old=confirm(client,items)
        new=client.post(f"/api/knowledge/documents/{doc['id']}/action",json={'action':'reextract'}).json()
        asyncio.run(app.state.knowledge_manager.process_run(new['run']['id']))
        items=extract(app,client,new)
        latest=[i for i in items if i['generation']==2]
        adopted=confirm(client,latest)
        assert adopted!=old
        client.post('/api/knowledge/review',json={'items':[{'id':latest[0]['id'],'status':'pending'}]})
        assert client.get('/api/knowledge/company').json()['revision']['fact_count']==0
        assert app.state.knowledge_store.revision(old)['fact_count']==1


def test_unexpected_batch_failure_requires_retry_and_never_becomes_complete(app):
    with TestClient(app) as client:
        doc=setup_document(app,client)
        class Broken:
            async def generate(self,*args,**kwargs):raise RuntimeError('internal error')
        extract(app,client,doc,Broken())
        state=app.state.knowledge_store.document(doc['id'])['run']
        assert state['status']=='failed' and state['failed_batches']==1 and state['completed_batches']==0
        response=client.post(f"/api/knowledge/documents/{doc['id']}/action",json={'action':'retry','expected_config':CONFIG})
        assert response.status_code==200
        app.state.ai_client=KnowledgeAI();asyncio.run(app.state.knowledge_manager.process_run(doc['run']['id']))
        assert app.state.knowledge_store.document(doc['id'])['run']['status']=='awaiting_review'


@pytest.mark.parametrize('changes',[{'kind':'material','name':'营业执照'},{'occurrence':'mention'},{'occurrence':'template'},{'subject':'制造商','occurrence':'third_party'},{'evidence_kind':'inference'},{'applicable_date':'2020-01-01'}])
def test_uncertain_and_historical_knowledge_cannot_support_hard_exclusion(app,changes):
    with TestClient(app) as client:
        doc=setup_document(app,client);batch=app.state.knowledge_store.next_batch(doc['run']['id'])
        output=ExtractionOutput.model_validate({'entries':[entry(**changes)],'limitations':[]})
        app.state.knowledge_store.finish_batch(doc['run']['id'],0,output,batch['units'])
        app.state.knowledge_store.update_run(doc['run']['id'],status='awaiting_review')
        if changes.get('occurrence') in {'mention','template'}:
            with app.state.knowledge_store.connect() as db:
                row=db.execute('SELECT id,source FROM entries WHERE run_id=?',(doc['run']['id'],)).fetchone()
                source=json.loads(row['source']);source.pop('reference_only',None)
                db.execute('UPDATE entries SET source=? WHERE id=?',(json.dumps(source),row['id']))
            assert app.state.knowledge_store.entries()['total']==0
            assert client.post('/api/knowledge/review',json={'items':[{'id':row['id'],'status':'accepted'}]}).status_code==422
            return
        snapshot=app.state.knowledge_store.revision(confirm(client,app.state.knowledge_store.entries()['items']))
        fact=snapshot['profile']['facts'][0]
        result=validate_decision(BidDecision.model_validate({'verdict':'exclude','summary':'模型误判','pending_checks':[],'requirements':[{'requirement':'民事责任能力','mandatory':True,'announcement_quote':'满足政府采购法第二十二条。','assessment':'conflict','company_fact_ids':[fact['id']],'company_quote':fact['quote'],'reason':'误用历史材料'}]}),CompanyProfile.model_validate(snapshot['profile']),ANNOUNCEMENT,snapshot['unsafe_fact_ids'])
        assert result.verdict=='needs_review' and result.requirements[0].assessment=='unknown'
