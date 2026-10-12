"""Local text extraction and image preparation; never silently drop pages."""
from __future__ import annotations

import base64
import io
import os
import shutil
import subprocess
import signal
import tempfile
import time
import math
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps

from .bid_models import BidCheckError

ALLOWED_EXTENSIONS = {"pdf", "doc", "docx", "xls", "xlsx", "png", "jpg", "jpeg", "txt"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TOTAL_BYTES = 50 * 1024 * 1024


def soffice_path() -> str | None:
    configured = os.getenv("CCGP_SOFFICE_PATH")
    if configured and Path(configured).is_file():
        return configured
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    bundled = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
    return str(bundled) if bundled.is_file() else None


def _convert(path: Path, extension: str, work: Path, *, check=None, progress=None) -> Path:
    check = check or (lambda: None)
    check()
    executable = soffice_path()
    if not executable:
        raise BidCheckError("Office 转换程序不可用，请安装 LibreOffice 或设置 CCGP_SOFFICE_PATH")
    work.mkdir(parents=True, mode=0o700, exist_ok=True)
    output = Path(tempfile.mkdtemp(prefix="office-convert-", dir=work))
    profile = output / "office-profile"
    try:
        timeout_seconds = float(os.getenv("CCGP_OFFICE_CONVERT_TIMEOUT_SECONDS", "1200"))
        if not math.isfinite(timeout_seconds): timeout_seconds = 1200
    except ValueError:
        timeout_seconds = 1200
    timeout_seconds = max(60, min(3600, timeout_seconds))
    process = None
    def terminate():
        if process is None: return
        try: os.killpg(process.pid, signal.SIGTERM)
        except OSError: pass
        try: process.communicate(timeout=2)
        except subprocess.TimeoutExpired:
            try: os.killpg(process.pid, signal.SIGKILL)
            except OSError: pass
            process.communicate()
    try:
        try:
            process = subprocess.Popen(
                [executable, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless", "--convert-to", extension, "--outdir", str(output), str(path)],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True,
            )
        except OSError:
            raise BidCheckError("Office 转换程序无法启动，请检查 LibreOffice 安装") from None
        deadline = time.monotonic() + timeout_seconds
        last_report = 0
        while True:
            check()
            current = time.monotonic()
            if progress and current-last_report>=1:
                progress();last_report=current
            remaining = deadline - current
            if remaining <= 0:
                raise BidCheckError(f"Office 转换超时（{int(timeout_seconds)} 秒），原件已保留，可重新读取")
            try:
                stdout, stderr = process.communicate(timeout=min(0.5, remaining))
                break
            except subprocess.TimeoutExpired: pass
        check()
        converted = output / f"{path.stem}.{extension}"
        if process.returncode:
            raise BidCheckError(f"Office 转换退出失败（退出码 {process.returncode}），请检查原文件")
        if not converted.is_file() or not converted.stat().st_size:
            raise BidCheckError("Office 未生成有效文件，原件可能损坏或格式不受支持")
        shutil.rmtree(profile, ignore_errors=True)
        os.chmod(converted, 0o600)
        return converted
    except BaseException:
        terminate()
        shutil.rmtree(output, ignore_errors=True)
        raise


def extract_documents(files: list[dict], directory: Path) -> tuple[list[dict], list[dict]]:
    sources: list[dict] = []
    images: list[dict] = []

    def add_text(source_id: str, text: str):
        if text.strip():
            sources.append({"source_id": source_id, "kind": "text", "text": text.strip()})
            if sum(len(source["text"]) for source in sources) > 200_000:
                raise BidCheckError("公司资料文字超出单次分析范围，请精简后上传；未截断资料")

    def add_image(source_id: str, content: bytes):
        try:
            with Image.open(io.BytesIO(content)) as raw:
                img = ImageOps.exif_transpose(raw).convert("RGB")
                if img.width * img.height > 40_000_000:
                    raise BidCheckError("图片分辨率过大，请缩小图片后重新上传")
                img.thumbnail((2400, 2400))
                output = io.BytesIO()
                img.save(output, format="JPEG", quality=90)
        except BidCheckError:
            raise
        except Exception:
            raise BidCheckError("公司图片无法读取，请确认文件没有损坏") from None
        sources.append({"source_id": source_id, "kind": "image", "text": ""})
        images.append({"source_id": source_id, "data_url": "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode()})
        if len(images) > 100 or sum(len(image["data_url"]) for image in images) > MAX_TOTAL_BYTES:
            raise BidCheckError("公司资料图像超出单次分析范围，请拆分后上传；未跳过页面")

    def add_pdf(pdf_path: Path, label: str, *, render_all=False):
        with pymupdf.open(pdf_path) as doc:
            if doc.is_encrypted:
                raise BidCheckError("PDF 已加密，请上传可读取的文件")
            for number, page in enumerate(doc, 1):
                page_label = f"{label}:第{number}页"
                text = page.get_text()
                add_text(page_label, text)
                if render_all or len(text.strip()) < 50 or page.get_images():
                    add_image(page_label + ":图像", page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png"))

    for index, file in enumerate(files, 1):
        path = directory / file["stored_name"]
        label = f"F{index}:{file['name']}"
        extension = path.suffix.lower()[1:]
        before = len(sources)
        try:
            if extension == "doc":
                path = _convert(path, "pdf", directory)
                extension = "pdf"
            elif extension == "xls":
                path = _convert(path, "xlsx", directory)
                extension = "xlsx"
            if extension == "txt":
                raw = path.read_bytes()
                try:
                    text = raw.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = raw.decode("gb18030")
                add_text(label, text)
            elif extension == "pdf":
                add_pdf(path, label, render_all=file["stored_name"].endswith(".doc"))
            elif extension == "docx":
                doc = Document(path)
                add_text(label + ":正文", doc.element.body.xml and "\n".join(node.text or "" for node in doc.element.body.iter() if node.tag.endswith("}t")))
                for section_index, section in enumerate(doc.sections, 1):
                    for kind, part in (("页眉", section.header), ("页脚", section.footer)):
                        add_text(label + f":第{section_index}节{kind}", "\n".join(node.text or "" for node in part._element.iter() if node.tag.endswith("}t")))
                with zipfile.ZipFile(path) as archive:
                    requires_render = any(name.startswith("word/charts/") for name in archive.namelist())
                    for name in archive.namelist():
                        if name.startswith("word/media/"):
                            if Path(name).suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff"}:
                                requires_render = True
                                continue
                            if archive.getinfo(name).file_size > MAX_FILE_BYTES:
                                raise BidCheckError("文档内图片过大，请拆分资料后上传")
                            add_image(label + f":嵌入图{Path(name).name}", archive.read(name))
                    if requires_render:
                        add_pdf(_convert(path, "pdf", directory), label + ":完整视觉页", render_all=True)
            elif extension == "xlsx":
                workbook = load_workbook(path, data_only=True)
                formulas = load_workbook(path, data_only=False, read_only=True)
                try:
                    has_charts = False
                    for sheet in workbook:
                        has_charts = has_charts or bool(sheet._charts)
                        rows = []
                        for row in sheet.iter_rows():
                            for cell in row:
                                if cell.value is None and formulas[sheet.title][cell.coordinate].data_type == "f":
                                    raise BidCheckError("Excel 包含未计算的公式，请在表格软件中计算保存或转换 PDF 后上传")
                            cells = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                            if cells:
                                rows.append(" | ".join(cells))
                        add_text(label + f":工作表{sheet.title}", "\n".join(rows))
                        for number, image in enumerate(sheet._images, 1):
                            add_image(label + f":工作表{sheet.title}:图{number}", image._data())
                    if has_charts:
                        add_pdf(_convert(path, "pdf", directory), label + ":图表视觉页", render_all=True)
                finally:
                    workbook.close()
                    formulas.close()
            else:
                add_image(label, path.read_bytes())
        except BidCheckError:
            raise
        except Exception:
            raise BidCheckError(f"第 {index} 份公司资料无法完整读取，请检查格式、内容或拆分后重新上传") from None
        if len(sources) == before:
            raise BidCheckError(f"第 {index} 份公司资料没有可读取的内容，请补充资料")
    total_text = sum(len(source["text"]) for source in sources)
    total_images = sum(len(image["data_url"]) for image in images)
    if total_text > 200_000 or total_images > MAX_TOTAL_BYTES or len(images) > 100:
        raise BidCheckError("公司资料超出单次分析范围，请拆分或精简后重新上传；未截断资料")
    return sources, images
