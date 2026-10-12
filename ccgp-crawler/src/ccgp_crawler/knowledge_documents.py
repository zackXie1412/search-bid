"""Index long documents without retaining standalone images; render batches on demand."""
from __future__ import annotations
import base64
import io
import time
import warnings
import zipfile
from pathlib import Path
import pymupdf
from PIL import Image, ImageOps
from openpyxl import load_workbook
from .bid_documents import _convert
from .docx_reader import DocxReader, NeedsOfficeRendering
from .bid_models import BidCheckError

MAX_PAGES = 3000
BATCH_CHARS = 20000
BATCH_UNITS = 6
MAX_IMAGE_BATCH = 15 * 1024 * 1024
DOCX_BLOCK_CHARS = 4000
DOCX_MAX_IMAGE_PIXELS = 100_000_000

def index_document(path: Path, name: str, work: Path, *, progress=None, check=None):
    progress = progress or (lambda **kw: None)
    check = check or (lambda: None)
    started = time.monotonic()
    stats = {}
    def report(**values):
        progress(elapsed_ms=round((time.monotonic()-started)*1000), **values)
    def convert(source, target):
        stats["mode"]="office_convert"
        report(phase="office_convert", **stats)
        return _convert(source, target, work, check=check, progress=lambda:report(phase="office_convert", **stats))
    units = []
    extension = path.suffix.lower()
    def add(location, text, visual=None):
        # Split large text blocks, retaining every character and the same source location.
        parts = [text[i:i+BATCH_CHARS] for i in range(0, len(text), BATCH_CHARS)] or ['']
        for index, part in enumerate(parts):
            units.append({'source_id': f'S{len(units)+1}', 'location': location + (f' · 片段{index+1}' if len(parts)>1 else ''), 'text': part, 'visual': visual if index == 0 else None})
    try:
        if extension == '.docx':
            try:
                reader = DocxReader(path, progress=report, check=check)
                units, stats = reader.read()
                pages = len(units)
                if pages > MAX_PAGES: raise BidCheckError('DOCX 超过 3,000 个内容块，请拆分上传；未截断内容')
            except NeedsOfficeRendering as exc:
                stats = {**dict(reader.stats),'unreadable': list(dict.fromkeys(reader.unsupported)), 'fallback_reason': str(exc)}
                units = []
                path = convert(path, 'pdf')
                extension = '.pdf'
            else:
                extension = '.docx_direct'
        elif extension == '.doc':
            path = convert(path, 'pdf')
            extension = '.pdf'
        if extension == '.pdf':
            with pymupdf.open(path) as pdf:
                if pdf.is_encrypted:
                    raise BidCheckError('PDF 已加密，请上传可读取的原件')
                if len(pdf) > MAX_PAGES:
                    raise BidCheckError('文档超过 3,000 页，请拆分上传；未跳过页面')
                for number, page in enumerate(pdf):
                    check()
                    text = page.get_text()
                    # Full-page render preserves embedded evidence and scanned pages.
                    visual = {'kind':'pdf', 'path': str(path), 'page':number} if len(text.strip())<50 or page.get_images() or page.get_drawings() else None
                    add(f'第{number+1}页', text, visual)
                pages = len(pdf)
        elif extension in {'.xls', '.xlsx'}:
            if extension == '.xls':
                path = convert(path, 'xlsx')
            book = load_workbook(path, data_only=True)
            formulas = load_workbook(path, data_only=False, read_only=True)
            try:
                if len(book.sheetnames) > MAX_PAGES:
                    raise BidCheckError('工作表数量超过 3,000，请拆分上传')
                if any(sheet._charts for sheet in book):
                    return index_document(convert(path, 'pdf'), name, work, progress=progress, check=check)
                for sheet in book:
                    lines = []
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value is None and formulas[sheet.title][cell.coordinate].data_type == 'f':
                                raise BidCheckError('Excel 包含未计算的公式，请计算保存或转换 PDF 后上传')
                        cells = [f'{cell.coordinate}={cell.value}' for cell in row if cell.value is not None]
                        if cells: lines.append(' | '.join(cells))
                    add(f'工作表 {sheet.title}', '\n'.join(lines))
                    for index, img in enumerate(sheet._images):
                        add(f'工作表 {sheet.title} · 图{index+1}', '', {'kind':'xlsx', 'path':str(path), 'sheet':sheet.title, 'image':index})
                pages = len(book.sheetnames)
            finally:
                book.close(); formulas.close()
        elif extension == '.txt':
            data = path.read_bytes()
            try: text = data.decode('utf-8-sig')
            except UnicodeDecodeError: text = data.decode('gb18030')
            if not text.strip(): raise BidCheckError('文件没有可读取的内容')
            add('正文', text)
            pages = len(units)
            if pages > MAX_PAGES: raise BidCheckError('文档超过 3,000 个内容块，请拆分上传；未截断内容')
        elif extension != '.docx_direct':
            # Decode once during indexing to reject corrupt/oversized images.
            image_bytes({'kind':'image', 'path':str(path)})
            add('图片原件', '', {'kind':'image', 'path':str(path)})
            pages = 1
    except BidCheckError: raise
    except Exception:
        check()
        raise BidCheckError('原文件无法完整读取，请检查文件格式或转换为 PDF 后上传') from None
    if not units: raise BidCheckError('文档没有可读取的内容')
    stats = {**stats, 'mode': stats.get('mode', 'office_convert' if 'fallback_reason' in stats else extension.lstrip('.')), 'blocks':len(units)}
    report(phase='image_prepare', **stats, processed=0, total=len(units))
    batches, current, chars, image_size = [], [], 0, 0
    sizes = {}; unreadable = []
    for position, unit in enumerate(units, 1):
        check()
        size = 0
        if unit['visual']:
            key = tuple(sorted(unit['visual'].items()))
            try:
                if key not in sizes:
                    sizes[key] = 4*((len(image_bytes(unit['visual']))+2)//3)+23
                size = sizes[key]
                if size > MAX_IMAGE_BATCH: raise BidCheckError('单张图像超过 15 MB')
            except Exception as exc:
                unreadable.append(f"{unit['location']}：{exc if isinstance(exc, BidCheckError) else '图片损坏或无法解码'}")
        if current and (len(current)>=BATCH_UNITS or chars+len(unit['text'])>BATCH_CHARS or image_size+size>MAX_IMAGE_BATCH):
            batches.append(current);current=[];chars=0;image_size=0
        current.append(unit); chars += len(unit['text']);image_size+=size
        if position % 10 == 0 or position == len(units):
            report(phase='image_prepare', **stats, processed=position, total=len(units), unreadable_images=unreadable)
    check()
    if unreadable:
        report(phase='image_prepare', **{**stats,'unreadable':unreadable}, processed=len(units),total=len(units))
        raise BidCheckError('读取不完整：'+'；'.join(unreadable)[:1500])
    if current: batches.append(current)
    stats.update(batches=len(batches),units=len(units),processed=len(units),total=len(units))
    # Successful rendering covers structures the direct reader could not interpret.
    if stats.get('fallback_reason'):stats['unreadable']=[]
    report(phase='indexing_complete', **stats)
    return pages, batches


def image_bytes(visual):
    if visual['kind']=='docx':
        with zipfile.ZipFile(visual['path']) as archive:
            content = archive.read(visual['media'])
    elif visual['kind']=='pdf':
        with pymupdf.open(visual['path']) as pdf:
            content = pdf[visual['page']].get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).tobytes('png')
    elif visual['kind']=='xlsx':
        book=load_workbook(visual['path'])
        try: content=book[visual['sheet']]._images[visual['image']]._data()
        finally: book.close()
    else: content=Path(visual['path']).read_bytes()
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', Image.DecompressionBombWarning)
        try:return _resize_image(content)
        except Image.DecompressionBombError:
            raise BidCheckError('嵌入图片超过 1 亿像素，请缩小后上传') from None


def _resize_image(content):
    with Image.open(io.BytesIO(content)) as raw:
        if raw.width*raw.height>DOCX_MAX_IMAGE_PIXELS: raise BidCheckError('嵌入图片超过 1 亿像素，请缩小后上传')
        raw.thumbnail((2400,2400))
        image=ImageOps.exif_transpose(raw).convert('RGB')
        stream=io.BytesIO();image.save(stream,format='JPEG',quality=85)
        return stream.getvalue()


def batch_images(units):
    images=[];size=0
    for unit in units:
        if unit.get('visual'):
            data='data:image/jpeg;base64,'+base64.b64encode(image_bytes(unit['visual'])).decode()
            size+=len(data)
            if size>MAX_IMAGE_BATCH:
                raise BidCheckError('本批图像超过 15 MB，请降低原件图片分辨率后重新上传')
            images.append({'source_id':unit['source_id'], 'data_url':data})
    return images
